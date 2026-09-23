"""
run_research_backtest.py - Research Strategy Default Backtest Runner

Run a single MT5 backtest for a researched strategy and save results (charts +
Word report + summary manifest) to:
    researched_strategies/<EA_NAME>/

Inside researched_strategies/<EA_NAME>/:
  - Strategy logic word doc (*.docx)
  - EA MQL5 source code (*.mq5)
  - Default backtest results (HTML report, equity charts, Word report, summary JSON)

Usage:
    1. Edit config.json or pass environment variables.
    2. python run_research_backtest.py
"""

import os
import sys
import time
import json
import shutil
from pathlib import Path

# Prevent Windows console UnicodeEncodeError when running on cp1252 / charmap environments
try:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# ─────────────────────────────────────────────────────────────────────────────
#  Local Project & Configuration Resolution
# ─────────────────────────────────────────────────────────────────────────────

_SCRIPT_DIR = Path(__file__).parent.resolve()
_CONFIG_PATH = _SCRIPT_DIR / "config.json"

_cfg = {}
if _CONFIG_PATH.exists():
    try:
        with open(_CONFIG_PATH, "r", encoding="utf-8") as f:
            _cfg = json.load(f)
    except Exception as e:
        print(f"  [WARN] Failed to load config.json: {e}")

def _env(key, default):
    return os.environ.get(f"AF_{key}", default)

# Base MT5 and Account Settings
TERMINAL_PATH = _env("TERMINAL_PATH", _cfg.get("terminal_path", r"C:\Users\HP\AppData\Roaming\MetaTrader\terminal64.exe"))
TERMINAL_DATA_DIR = _env("TERMINAL_DATA_DIR", _cfg.get("terminal_data_dir", r"C:\Users\HP\AppData\Roaming\MetaQuotes\Terminal\CDE1ED2F37049DA2E508A3C44B675D09"))

try:
    LOGIN = int(_env("LOGIN", str(_cfg.get("login", "52909674"))))
except Exception:
    LOGIN = 52909674

PASSWORD = _env("PASSWORD", _cfg.get("password", "3F!@4rwo7wc02f"))
SERVER   = _env("SERVER",   _cfg.get("server", "ICMarketsKE-Demo"))
CURRENCY = _env("CURRENCY", _cfg.get("currency", "USD"))
LEVERAGE = _env("LEVERAGE", _cfg.get("leverage", "1:100"))

try:
    SINGLE_TEST_TIMEOUT = int(_env("SINGLE_TEST_TIMEOUT", str(_cfg.get("single_test_timeout", 1800))))
except Exception:
    SINGLE_TEST_TIMEOUT = 1800

# Strategy and Market Settings
ACTIVE_EA = _cfg.get("active_ea", "TRB")
DEFAULT_EXPERT = _cfg.get("expert", f"{ACTIVE_EA} V2.0.ex5")

STRATEGY_NAME = _env("STRATEGY_NAME", _env("EA_NAME", ACTIVE_EA))
EXPERT        = _env("EXPERT", DEFAULT_EXPERT)
SET_FILE      = _env("SET_FILE", "")
SYMBOL        = _env("SYMBOL", _cfg.get("symbol", "USDJPY Dukascopy"))
PERIOD        = _env("PERIOD", _cfg.get("period", "M15"))
FROM_DATE     = _env("FROM_DATE", _cfg.get("train_from", "2013.01.01"))
TO_DATE       = _env("TO_DATE", _cfg.get("holdout_to", "2026.07.03"))
RESEARCH_DEPOSIT = float(_env("DEPOSIT", str(_cfg.get("deposit", "2500"))))

# ─────────────────────────────────────────────────────────────────────────────
#  Project Imports (with graceful fallbacks)
# ─────────────────────────────────────────────────────────────────────────────

try:
    from mt5_runner import run_single_backtest
except ImportError as e:
    print(f"[ERROR] Could not import mt5_runner: {e}")
    raise

try:
    from report_analysis import analyze
except ImportError as e:
    print(f"[WARN] Could not import report_analysis: {e}")
    analyze = None

try:
    from run_full_backtest import (
        parse_mt5_html_for_deals,
        generate_equity_charts,
        generate_monthly_heatmap,
        calculate_derived_metrics,
    )
    _HAS_FULL_BT = True
except Exception:
    _HAS_FULL_BT = False

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False


# ─────────────────────────────────────────────────────────────────────────────
#  Helpers & Word Reporting
# ─────────────────────────────────────────────────────────────────────────────

def _fmt(val, prefix="$", suffix="", decimals=2):
    if val is None:
        return "N/A"
    try:
        return f"{prefix}{float(val):,.{decimals}f}{suffix}"
    except (TypeError, ValueError):
        return str(val)


def _create_research_word_report(report_dir: Path, strategy_name: str,
                                 curated: dict, derived: dict,
                                 from_date: str, to_date: str,
                                 deposit: float, expert: str, symbol: str, period: str):
    if not _HAS_DOCX:
        print("    [WARN] python-docx not installed — skipping Word report generation.")
        return None

    doc = docx.Document()

    # Title
    title = doc.add_heading("Research Strategy Default Backtest Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub_run = sub.add_run(f"Strategy / EA: {strategy_name}")
    sub_run.bold = True
    sub_run.font.size = Pt(14)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info_p = doc.add_paragraph(
        f"Expert: {expert}  |  Symbol: {symbol}  |  Timeframe: {period}\n"
        f"Backtest Period: {from_date}  ->  {to_date}  |  Initial Deposit: ${deposit:,.2f}\n"
        f"Generated: {time.strftime('%Y-%m-%d %H:%M:%S UTC', time.gmtime())}"
    )
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(
        "This report documents the DEFAULT (un-optimised baseline) backtest for this researched strategy. "
        "It establishes baseline performance and trade mechanics directly from the EA's core MQL5 implementation "
        "prior to running genetic parameter optimization and robustness stress-tests."
    )

    # Key metrics table
    doc.add_heading("Key Performance Metrics", level=1)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"

    d = {**curated, **derived}
    metrics_to_show = [
        ("Net Profit ($)",             d.get("Net Profit") or d.get("total_net_profit")),
        ("Net Profit (%)",             d.get("Net Profit %") or d.get("profit_pct")),
        ("Profit Factor",              d.get("Profit Factor") or d.get("profit_factor")),
        ("Sharpe Ratio",               d.get("Sharpe Ratio") or d.get("sharpe")),
        ("Max Balance Drawdown (%)",   d.get("Max Balance Drawdown (%)") or d.get("max_dd_pct")),
        ("Recovery Factor",            d.get("Recovery Factor") or d.get("recovery_factor")),
        ("Win Rate (%)",               d.get("Win Rate %") or d.get("win_rate")),
        ("Total Trades",               d.get("Total Trades") or d.get("total_trades")),
        ("Largest Profit Trade ($)",   d.get("Largest Profit Trade") or d.get("largest_profit")),
        ("Largest Loss Trade ($)",     d.get("Largest Loss Trade") or d.get("largest_loss")),
    ]
    for label, val in metrics_to_show:
        row = tbl.add_row().cells
        row[0].text = str(label)
        try:
            row[1].text = f"{float(val):,.4f}" if isinstance(val, float) else str(val or "N/A")
        except (TypeError, ValueError):
            row[1].text = str(val or "N/A")

    # Equity & Distribution Charts
    chart_files = list(report_dir.glob("chart_*.png"))
    if chart_files:
        doc.add_heading("Equity & Performance Charts", level=1)
        for ch in chart_files:
            try:
                doc.add_paragraph(ch.stem.replace("_", " ").title()).runs[0].bold = True
                doc.add_picture(str(ch), width=Inches(6.2))
                doc.add_paragraph()
            except Exception as e:
                print(f"    [WARN] Could not insert chart {ch.name}: {e}")

    doc_path = report_dir / f"{strategy_name}_Research_Report.docx"
    try:
        doc.save(doc_path)
        return doc_path
    except PermissionError:
        alt = report_dir / f"{strategy_name}_Research_Report_{int(time.time())}.docx"
        doc.save(alt)
        return alt


# ─────────────────────────────────────────────────────────────────────────────
#  EA & Folder Resolution
# ─────────────────────────────────────────────────────────────────────────────

def resolve_researched_folder(strategy_name: str, expert: str, research_dir_override: str = None) -> tuple[Path, str, str]:
    """
    Resolves the target EA folder inside 'researched_strategies'.
    Folder structure:
      <project_root>/researched_strategies/<EA_NAME>/
    Returns:
      (report_dir: Path, ea_folder_name: str, expert_file: str)
    """
    # 1. Base directory
    if research_dir_override and Path(research_dir_override).exists():
        base_dir = Path(research_dir_override)
    elif _cfg.get("research_dir") and Path(_cfg["research_dir"]).exists():
        base_dir = Path(_cfg["research_dir"])
    else:
        base_dir = _SCRIPT_DIR / "researched_strategies"
    base_dir.mkdir(parents=True, exist_ok=True)

    # 2. Extract clean EA name
    clean_expert = expert.strip()
    if clean_expert.lower().endswith(".ex5"):
        clean_ea = clean_expert[:-4]
    elif clean_expert.lower().endswith(".mq5"):
        clean_ea = clean_expert[:-4]
    else:
        clean_ea = clean_expert

    # If strategy_name is specified and not generic "My_Strategy_v1", prioritize it
    ea_folder_name = clean_ea or "Researched_EA"
    if strategy_name and strategy_name != "My_Strategy_v1":
        ea_folder_name = strategy_name
    elif (base_dir / clean_ea).is_dir():
        ea_folder_name = clean_ea

    # Check existing folders in researched_strategies
    if not (base_dir / ea_folder_name).is_dir():
        for sub in base_dir.iterdir():
            if sub.is_dir() and (sub.name.lower() == ea_folder_name.lower() or sub.name.lower() == clean_ea.lower()):
                ea_folder_name = sub.name
                break

    report_dir = base_dir / ea_folder_name
    report_dir.mkdir(parents=True, exist_ok=True)

    # Expert filename normalization
    expert_file = expert.strip()
    if not expert_file:
        expert_file = f"{ea_folder_name}.ex5"
    if not expert_file.lower().endswith(".ex5"):
        expert_file += ".ex5"

    return report_dir, ea_folder_name, expert_file


# ─────────────────────────────────────────────────────────────────────────────
#  Main Runner
# ─────────────────────────────────────────────────────────────────────────────

def run_research_backtest(
    strategy_name: str = STRATEGY_NAME,
    set_file: str = SET_FILE,
    expert: str = EXPERT,
    symbol: str = SYMBOL,
    period: str = PERIOD,
    from_date: str = FROM_DATE,
    to_date: str = TO_DATE,
    deposit: float = None,
    research_dir: str = None,
):
    deposit = deposit if deposit is not None else RESEARCH_DEPOSIT

    # 1. Resolve folder in researched_strategies
    report_dir, ea_name, expert_filename = resolve_researched_folder(strategy_name, expert, research_dir)

    print("\n" + "=" * 74)
    print(f"  ALPHA FORGE -- RESEARCH STRATEGY BACKTEST")
    print(f"  Strategy / EA Name : {ea_name}")
    print(f"  Target Folder      : {report_dir}")
    print(f"  Expert (.ex5)      : {expert_filename}")
    print(f"  Symbol / Period    : {symbol}  |  {period}")
    print(f"  Date Range         : {from_date}  ->  {to_date}")
    print(f"  Initial Deposit    : ${deposit:,.2f} {CURRENCY}")
    print("=" * 74 + "\n")

    # 2. Inspect EA assets inside researched_strategies/<EA_NAME>/
    doc_files = list(report_dir.glob("*.docx")) + list(report_dir.glob("*.doc"))
    mq5_files = list(report_dir.glob("*.mq5"))
    ex5_files = list(report_dir.glob("*.ex5"))

    if doc_files:
        print(f"  [DOC] Found Strategy Logic Doc : {doc_files[0].name}")
    else:
        print("  [INFO] No Strategy Logic Word doc found in folder (optional)")

    if mq5_files:
        print(f"  [CODE] Found EA MQL5 Source Code: {mq5_files[0].name}")
    else:
        print("  [INFO] No MQL5 source code file found in folder (optional)")

    # 3. Synchronize compiled EA (.ex5) to MT5 Experts directory if needed
    mt5_experts_dir = Path(TERMINAL_DATA_DIR) / "MQL5" / "Experts"
    mt5_experts_dir.mkdir(parents=True, exist_ok=True)
    target_expert_in_mt5 = mt5_experts_dir / expert_filename

    # If the .ex5 is in the researched_strategies folder but not yet in MT5 Experts, copy it!
    matching_local_ex5 = [f for f in ex5_files if f.name.lower() == expert_filename.lower()]
    if matching_local_ex5 and not target_expert_in_mt5.exists():
        try:
            shutil.copy2(matching_local_ex5[0], target_expert_in_mt5)
            print(f"  [OK] Synchronized EA binary to MT5 Experts: {target_expert_in_mt5.name}")
        except Exception as e:
            print(f"  [WARN] Failed to copy .ex5 into MT5 Experts: {e}")

    # 4. Handle .set file resolution (Default vs Custom)
    # MT5 Strategy Tester requires the set file to be in MQL5\Profiles\Tester\
    profiles_tester_dir = Path(TERMINAL_DATA_DIR) / "MQL5" / "Profiles" / "Tester"
    profiles_tester_dir.mkdir(parents=True, exist_ok=True)

    clean_set = set_file.strip() if set_file else ""
    active_set_file = None

    # Check if a custom set file was specified and exists
    if clean_set and clean_set not in ("my_strategy.set", "(Default EA Inputs)", "(none)", "default"):
        p_set = Path(clean_set)
        if p_set.is_file():
            shutil.copy2(p_set, profiles_tester_dir / p_set.name)
            active_set_file = p_set.name
        elif (report_dir / clean_set).is_file():
            shutil.copy2(report_dir / clean_set, profiles_tester_dir / clean_set)
            active_set_file = clean_set
        elif (profiles_tester_dir / clean_set).is_file():
            active_set_file = clean_set

    # If no set file was found/specified, check if there is any .set file in the EA's research folder
    if not active_set_file:
        local_sets = list(report_dir.glob("*.set"))
        if local_sets:
            shutil.copy2(local_sets[0], profiles_tester_dir / local_sets[0].name)
            active_set_file = local_sets[0].name
            print(f"  [*] Using detected set file from folder: {active_set_file}")

    # If still no set file, create a default parameter file in Profiles/Tester
    # This runs the EA with its built-in default inputs and guarantees MT5 won't fail with 'file not found'
    if not active_set_file:
        safe_ea_tag = ea_name.replace(" ", "_").replace(".", "_")
        default_set_name = f"research_{safe_ea_tag}_default.set"
        default_set_path = profiles_tester_dir / default_set_name
        default_set_path.write_text(
            f"; AlphaForge Default Research Parameters for {ea_name}\n"
            f"; Testing EA default parameters from MQL5 code\n",
            encoding="utf-8"
        )
        active_set_file = default_set_name
        print(f"  [*] Using EA default built-in inputs ({active_set_file})")

    # 5. Execute MT5 Backtest
    report_name = f"{ea_name}_default"
    print(f"\n  [>] Launching MetaTrader 5 Strategy Tester...")
    print(f"    Terminal : {TERMINAL_PATH}")
    print(f"    Data Dir : {TERMINAL_DATA_DIR}")
    print(f"    Set File : {active_set_file}")

    try:
        report_path = run_single_backtest(
            terminal_path=TERMINAL_PATH,
            terminal_data_dir=TERMINAL_DATA_DIR,
            expert=expert_filename,
            set_file=active_set_file,
            symbol=symbol,
            period=period,
            from_date=from_date,
            to_date=to_date,
            work_dir=str(report_dir),
            login=LOGIN,
            password=PASSWORD,
            server=SERVER,
            report_name=report_name,
            deposit=int(deposit),
            currency=CURRENCY,
            leverage=LEVERAGE,
            timeout=SINGLE_TEST_TIMEOUT,
        )
    except Exception as exc:
        print(f"\n[ERROR] MT5 Backtest execution failed: {exc}")
        raise

    print(f"\n  [OK] Backtest finished! Report generated: {report_path.name}")

    # 6. Parse and Analyze Results
    curated = {}
    if analyze:
        try:
            result = analyze(report_path)
            curated = result.get("curated_summary", {})
        except Exception as e:
            print(f"  [WARN] Analysis parser error: {e}")

    derived = {}
    # Generate Equity Charts and Derived Metrics
    if _HAS_FULL_BT:
        try:
            deals_df = parse_mt5_html_for_deals(report_path, deposit=float(deposit))
            if deals_df is not None and not deals_df.empty:
                chart_paths = generate_equity_charts(deals_df, float(deposit), report_dir)
                derived = calculate_derived_metrics(deals_df, float(deposit))
                print(f"  [OK] Equity charts generated: {list(chart_paths.keys())}")
        except Exception as e:
            print(f"  [WARN] Could not generate advanced charts: {e}")

    # 7. Generate Word Report inside researched_strategies/<EA_NAME>/
    doc_path = _create_research_word_report(
        report_dir=report_dir,
        strategy_name=ea_name,
        curated=curated,
        derived=derived,
        from_date=from_date,
        to_date=to_date,
        deposit=deposit,
        expert=expert_filename,
        symbol=symbol,
        period=period,
    )

    # 8. Save Summary JSON Manifest
    net_p  = curated.get("Net Profit") or derived.get("total_net_profit", 0) or 0
    pct    = curated.get("Net Profit %") or derived.get("profit_pct", 0) or 0
    pf     = curated.get("Profit Factor") or derived.get("profit_factor", 0) or 0
    sharpe = curated.get("Sharpe Ratio") or derived.get("sharpe", 0) or 0
    maxdd  = curated.get("Max Balance Drawdown (%)") or derived.get("max_dd_pct", 0) or 0
    trades = curated.get("Total Trades") or derived.get("total_trades", 0) or 0
    wr     = curated.get("Win Rate %") or derived.get("win_rate", 0) or 0

    manifest = {
        "ea_name": ea_name,
        "expert": expert_filename,
        "symbol": symbol,
        "period": period,
        "date_range": f"{from_date} -> {to_date}",
        "deposit": deposit,
        "strategy_doc": doc_files[0].name if doc_files else None,
        "mql5_code": mq5_files[0].name if mq5_files else None,
        "word_report": doc_path.name if doc_path else None,
        "html_report": report_path.name,
        "metrics": {
            "net_profit": float(net_p),
            "profit_pct": float(pct),
            "profit_factor": float(pf),
            "sharpe_ratio": float(sharpe),
            "max_drawdown_pct": float(maxdd),
            "total_trades": int(trades),
            "win_rate_pct": float(wr),
        },
        "timestamp_utc": time.strftime("%Y-%m-%d %H:%M:%S UTC", time.gmtime())
    }
    manifest_path = report_dir / f"{ea_name}_summary.json"
    with open(manifest_path, "w", encoding="utf-8") as mf:
        json.dump(manifest, mf, indent=2)

    # 9. Print Terminal Summary
    print("\n" + "-" * 74)
    print(f"  RESEARCH BACKTEST SUMMARY -- {ea_name}")
    print("-" * 74)
    print(f"  Net Profit      : ${float(net_p):>12,.2f}  ({float(pct):.2f}%)")
    print(f"  Profit Factor   : {float(pf):>10.3f}")
    print(f"  Sharpe Ratio    : {float(sharpe):>10.3f}")
    print(f"  Max Drawdown    : {float(maxdd):>9.2f}%")
    print(f"  Total Trades    : {int(trades):>10,}")
    print(f"  Win Rate        : {float(wr):>9.2f}%")
    print("-" * 74)
    print(f"  Results saved to: {report_dir}")
    if doc_path:
        print(f"  Word Report     : {doc_path.name}")
    print(f"  HTML Report     : {report_path.name}")
    print("=" * 74 + "\n")

    return {
        "report_path": report_path,
        "report_dir": report_dir,
        "curated": curated,
        "derived": derived,
        "doc_path": doc_path,
        "manifest": manifest,
    }


if __name__ == "__main__":
    run_research_backtest()
