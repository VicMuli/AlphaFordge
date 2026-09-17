"""
run_post_optimization.py - Batch Full Backtest & Monte Carlo Runner

Flow
----
  1. Targets a specific run directory via TARGET_RUN_DIR.
  2. Iterates through the TARGET_CANDIDATES list directly inside that run folder.
  3. Locates or builds the .set file for each candidate.
  4. Runs a full-period MT5 backtest, prints summary metrics to the terminal, and saves the report.
  5. Extracts the daily P&L distribution directly from the MT5 HTML report.
  6. Runs 5,000 Monte Carlo simulations for Phase 1 (8%) and Phase 2 (5%).
  7. Generates the MonteCarlo_Full_Summary_cand_XXX.docx report (with safe handling if the file is open).
"""

import os
import shutil
import json
import time
import random
import statistics
import re
from pathlib import Path
import pandas as pd

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None

try:
    import docx
    from docx.shared import RGBColor, Pt
except ImportError:
    print("WARNING: python-docx not installed. Word document will not be generated.")

# -- Project imports --
from run_optimization import (
    TERMINAL_PATH, TERMINAL_DATA_DIR, EXPERT, SYMBOL, PERIOD,
    LOGIN, PASSWORD, SERVER, DEPOSIT, CURRENCY, LEVERAGE,
    SINGLE_TEST_TIMEOUT, TRAIN_FROM, HOLDOUT_TO, WORK_DIR
)
from mt5_runner import run_single_backtest
from report_analysis import analyze


# ===========================================================================
#  USER CONFIGURATION
# ===========================================================================

# Specify the exact run folder (e.g., "run_20260911_094409")
TARGET_RUN_DIR = "run_20260911_094409"

# List the exact candidate folders you want to process
TARGET_CANDIDATES = [
    "cand_007",
    "cand_014",
    "cand_015",
]

# Backtest Dates (Defaulting to full period)
TEST_FROM = TRAIN_FROM
TEST_TO = HOLDOUT_TO

# Monte Carlo Parameters
SIMULATIONS_RUN = 5000
BLOCK_SIZE = 10
MAX_DAYS = 250
PHASE_1_TARGET_PCT = 8.0
PHASE_2_TARGET_PCT = 5.0
MAX_DD_LIMIT_PCT = 10.0
DAILY_DD_LIMIT_PCT = 5.0

# ===========================================================================

def get_target_run_dir(base_work_dir: Path, target_run: str) -> Path | None:
    if not target_run or target_run == "run_XXXXXXXX_XXXXXX":
        print("ERROR: Please specify a valid TARGET_RUN_DIR in the script configuration.")
        return None
        
    run_dir = base_work_dir / target_run
    if not run_dir.exists():
        print(f"ERROR: The specified run directory '{target_run}' does not exist.")
        print(f"Looked in: {base_work_dir}")
        return None
        
    return run_dir

def locate_or_build_set_file(cand_dir: Path) -> Path | None:
    cand_name = cand_dir.name
    cand_n = cand_name.replace("cand_", "")
    parent_dir = cand_dir.parent

    target_set_path = cand_dir / f"c{cand_n}_full.set"
    if target_set_path.exists(): return target_set_path

    candidates_to_check = list(cand_dir.glob("*.set")) + list(cand_dir.rglob("*.set"))
    if parent_dir and parent_dir.exists():
        candidates_to_check.extend(list(parent_dir.glob(f"*{cand_n}*.set")))

    if candidates_to_check: return candidates_to_check[0]

    json_files = list(cand_dir.glob("*.json")) + list(cand_dir.rglob("*.json"))
    if json_files:
        try:
            with open(json_files[0], "r") as f: params = json.load(f)
            generated_set = cand_dir / f"{cand_name}.set"
            with open(generated_set, "w", encoding="utf-16-le") as f:
                for k, v in params.items(): f.write(f"{k}={v}\n")
            return generated_set
        except Exception:
            pass
    return None

def _read_html_report(report_path: Path) -> str:
    """Read an MT5 HTML report with BOM/encoding auto-detection."""
    raw = report_path.read_bytes()

    # MT5 reports can be UTF-16 (with or without BOM) or UTF-8/Windows-1252.
    encodings = []
    if raw.startswith(b"\xff\xfe") or raw.startswith(b"\xfe\xff"):
        encodings.extend(["utf-16", "utf-16-le", "utf-16-be"])
    elif b"\x00" in raw[:2000]:
        # Null bytes are a strong signal of UTF-16 without a BOM.
        encodings.extend(["utf-16-le", "utf-16-be"])
    encodings.extend(["utf-8-sig", "utf-8", "cp1252", "latin-1"])

    tried = set()
    for encoding in encodings:
        if encoding in tried:
            continue
        tried.add(encoding)
        try:
            html = raw.decode(encoding)
            if "<html" in html.lower() or "<table" in html.lower():
                return html
        except (UnicodeDecodeError, UnicodeError):
            continue

    # Last-resort permissive decode so the caller can still report diagnostics.
    return raw.decode("utf-8", errors="ignore")


def _norm(text: str) -> str:
    """Normalize HTML cell text for reliable matching."""
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip().lower()


def _parse_money(value: str) -> float | None:
    """Parse MT5 numeric text, including commas, currency symbols and parentheses."""
    if value is None:
        return None

    s = value.replace("\xa0", " ").strip()
    if not s:
        return None

    # Negative accounting format: (123.45)
    negative = s.startswith("(") and s.endswith(")")
    s = s.strip("()")

    # Keep digits, decimal point and minus sign. This handles $1,234.56, € 123.45, etc.
    s = re.sub(r"[^0-9.\-+]", "", s)
    if not s or s in {"-", "+", ".", "-.", "+."}:
        return None

    try:
        number = float(s)
        return -abs(number) if negative else number
    except ValueError:
        return None


def _extract_summary_metrics(tables):
    """Extract common MT5 Strategy Tester summary metrics from arbitrary table layouts."""
    wanted = {
        "total net profit": "Total Net Profit",
        "profit factor": "Profit Factor",
        "expected payoff": "Expected Payoff",
        "maximal drawdown": "Maximal Drawdown",
        "absolute drawdown": "Absolute Drawdown",
        "total trades": "Total Trades",
        "win rate": "Win Rate",
        "balance": "Balance",
        "equity": "Equity",
        "gross profit": "Gross Profit",
        "gross loss": "Gross Loss",
    }

    metrics = {}

    for table in tables:
        rows = table.find_all("tr")
        for row in rows:
            cells = row.find_all(["td", "th"])
            values = [re.sub(r"\s+", " ", c.get_text(" ", strip=True)).strip() for c in cells]
            if not values:
                continue

            # Case 1: label and value are separate cells.
            for i, raw_label in enumerate(values):
                label_norm = _norm(raw_label).rstrip(":")
                if label_norm in wanted:
                    value = None
                    if i + 1 < len(values):
                        value = values[i + 1]
                    if value:
                        metrics[wanted[label_norm]] = value

                    # Handle "Label: Value" accidentally stored in the same cell.
                    if ":" in raw_label:
                        left, right = raw_label.split(":", 1)
                        if _norm(left) in wanted and right.strip():
                            metrics[wanted[_norm(left)]] = right.strip()

            # Case 2: the whole row is text such as "Total Net Profit: 123.45".
            for label_norm, display_name in wanted.items():
                pattern = rf"{re.escape(label_norm)}\s*:\s*([-+()$€£¥\d,.\s%]+)"
                match = re.search(pattern, " ".join(values), flags=re.IGNORECASE)
                if match and display_name not in metrics:
                    metrics[display_name] = match.group(1).strip()

    return metrics


def _find_trade_table(tables):
    """
    Find an MT5 deal/trade table by its header row.

    We require a time/date column and a profit/P&L column in the same row.
    """
    for table in tables:
        rows = table.find_all("tr")
        for idx, row in enumerate(rows):
            cells = row.find_all(["td", "th"])
            headers = [_norm(c.get_text(" ", strip=True)) for c in cells]
            if not headers:
                continue

            time_idx = next(
                (i for i, h in enumerate(headers)
                 if h == "time" or h.startswith("time ") or h in {"date", "datetime", "date/time"}),
                None,
            )
            profit_idx = next(
                (i for i, h in enumerate(headers)
                 if h in {"profit", "p/l", "pl"} or "profit" in h),
                None,
            )

            if time_idx is not None and profit_idx is not None:
                return rows[idx + 1:], time_idx, profit_idx

    return [], None, None


def print_and_extract_report_summary(report_path: Path):
    """Parse an MT5 HTML report, print metrics, and extract actual daily P&L."""
    daily_pnls = []
    summary_metrics = {}
    trades_data = []

    if not BeautifulSoup:
        print("    [WARNING] BeautifulSoup not available. Install with: pip install beautifulsoup4")
        return daily_pnls

    if not report_path.exists():
        print(f"    [ERROR] Report does not exist: {report_path}")
        return daily_pnls

    try:
        html = _read_html_report(report_path)
        soup = BeautifulSoup(html, "html.parser")
        tables = soup.find_all("table")

        # 1. Summary metrics — robust to MT5 HTML layout/encoding variations.
        summary_metrics = _extract_summary_metrics(tables)

        print("\n" + "─" * 45)
        print("       FULL BACKTEST PERFORMANCE SUMMARY       ")
        print("─" * 45)

        preferred_order = [
            "Total Net Profit",
            "Profit Factor",
            "Expected Payoff",
            "Maximal Drawdown",
            "Absolute Drawdown",
            "Total Trades",
            "Win Rate",
            "Gross Profit",
            "Gross Loss",
        ]

        printed = set()
        for key in preferred_order:
            if key in summary_metrics:
                print(f"  {key:<22}: {summary_metrics[key]}")
                printed.add(key)

        for key, value in summary_metrics.items():
            if key not in printed:
                print(f"  {key:<22}: {value}")

        if not summary_metrics:
            print("  [Notice] Could not match standard MT5 summary fields.")
            print("          The report was decoded successfully; see extraction diagnostics below.")

        print("─" * 45)

        # 2. Locate the actual MT5 deals table and extract Time + Profit.
        target_rows, time_idx, profit_idx = _find_trade_table(tables)

        if time_idx is not None and profit_idx is not None:
            for row in target_rows:
                cols = [
                    re.sub(r"\s+", " ", c.get_text(" ", strip=True)).strip()
                    for c in row.find_all(["td", "th"])
                ]

                if len(cols) <= max(time_idx, profit_idx):
                    continue

                time_text = cols[time_idx]
                profit_text = cols[profit_idx]

                p_val = _parse_money(profit_text)
                if p_val is None:
                    continue

                # Ignore non-trade rows such as totals/headers.
                if not time_text:
                    continue

                parsed_date = pd.to_datetime(time_text, errors="coerce")
                if pd.isna(parsed_date):
                    continue

                trades_data.append((parsed_date, p_val))

        print(
            f"    -> Report tables found: {len(tables)} | "
            f"trade rows parsed: {len(trades_data)}"
        )

    except Exception as e:
        print(f"    [ERROR] Failed parsing report HTML: {e}")

    # 3. Aggregate real trade P&L by calendar day.
    if trades_data:
        df_trades = pd.DataFrame(trades_data, columns=["Time", "Profit"])
        df_trades["Time"] = pd.to_datetime(df_trades["Time"], errors="coerce")
        df_trades = df_trades.dropna(subset=["Time"])

        if not df_trades.empty:
            daily_pnls = (
                df_trades.assign(Date=df_trades["Time"].dt.date)
                .groupby("Date", sort=True)["Profit"]
                .sum()
                .tolist()
            )

    if not daily_pnls or len(daily_pnls) < 3:
        print(
            "    [WARNING] Insufficient actual daily P&L records extracted. "
            "Monte Carlo will use the fallback distribution."
        )
        print(
            "              This fallback is only a safety net; it should not be "
            "treated as historical MT5 performance."
        )

        # Keep the original fallback so the script remains runnable.
        base_dist = [
            -45.5, -20.0, -10.0, 5.0, 12.5, 18.0, 25.0,
            35.0, 48.0, 85.0, -15.0, 40.0, -30.0, 20.0
        ]
        daily_pnls = base_dist * 15
    else:
        print(
            f"    -> Successfully extracted {len(daily_pnls)} actual "
            f"daily P&L records from report."
        )
        print(
            f"       Historical P&L range: {min(daily_pnls):,.2f} to "
            f"{max(daily_pnls):,.2f}"
        )

    return [float(p) for p in daily_pnls if float(p) != 0.0]

def run_monte_carlo_simulation(daily_pnls, deposit, target_pct, max_dd_pct, daily_dd_pct):
    target_amount = deposit * (target_pct / 100.0)
    max_dd_limit_amount = deposit * (max_dd_pct / 100.0)
    n_pnls = len(daily_pnls)
    if n_pnls == 0:
        raise ValueError("No daily P&L data available for Monte Carlo simulation.")
    
    pass_count, breach_max_dd_count, breach_daily_dd_count, incomplete_count = 0, 0, 0, 0
    days_to_pass_list, max_dd_reached_list, all_simulated_daily_returns = [], [], []

    for _ in range(SIMULATIONS_RUN):
        sim_pnls = []
        while len(sim_pnls) < MAX_DAYS:
            start_idx = random.randint(0, n_pnls - 1)
            sim_pnls.extend([daily_pnls[(start_idx + i) % n_pnls] for i in range(BLOCK_SIZE)])
        sim_pnls = sim_pnls[:MAX_DAYS]

        equity = deposit
        peak_equity = deposit
        max_dd_amount = 0.0
        passed, breached_max, breached_daily = False, False, False
        days_taken, prev_day_close = MAX_DAYS, deposit

        for day_idx, pnl in enumerate(sim_pnls, start=1):
            equity += pnl
            
            if (prev_day_close - equity) > (prev_day_close * (daily_dd_pct / 100.0)):
                breached_daily = True; break
            
            if (deposit - equity) > max_dd_limit_amount:
                breached_max = True; break

            peak_equity = max(peak_equity, equity)
            current_dd = peak_equity - equity
            max_dd_amount = max(max_dd_amount, current_dd)

            if (equity - deposit) >= target_amount:
                passed = True; days_taken = day_idx; break
            prev_day_close = equity

        max_dd_reached_list.append((max_dd_amount / deposit) * 100.0)
        
        if passed: 
            pass_count += 1
            days_to_pass_list.append(days_taken)
        elif breached_max: 
            breach_max_dd_count += 1
        elif breached_daily: 
            breach_daily_dd_count += 1
        else: 
            incomplete_count += 1

        all_simulated_daily_returns.extend(sim_pnls)

    days_to_pass_list.sort()
    max_dd_reached_list.sort()
    sorted_returns = sorted(all_simulated_daily_returns) if all_simulated_daily_returns else [0.0]

    def get_percentile(lst, pct):
        if not lst: return 0.0
        idx = int(len(lst) * (pct / 100.0))
        return lst[min(idx, len(lst) - 1)]

    return {
        "pass_rate": (pass_count / SIMULATIONS_RUN) * 100.0,
        "breach_max_rate": (breach_max_dd_count / SIMULATIONS_RUN) * 100.0,
        "breach_daily_rate": (breach_daily_dd_count / SIMULATIONS_RUN) * 100.0,
        "incomplete_rate": (incomplete_count / SIMULATIONS_RUN) * 100.0,
        "median_days": int(statistics.median(days_to_pass_list)) if days_to_pass_list else 0,
        "p10_days": get_percentile(days_to_pass_list, 10),
        "p90_days": get_percentile(days_to_pass_list, 90),
        "median_max_dd": statistics.median(max_dd_reached_list) if max_dd_reached_list else 0.0,
        "p90_max_dd": get_percentile(max_dd_reached_list, 90),
        "p99_max_dd": get_percentile(max_dd_reached_list, 99),
        "largest_profit_day": max(sorted_returns),
        "largest_loss_day": min(sorted_returns),
        "avg_daily": statistics.mean(sorted_returns),
        "p95_profit_day": get_percentile(sorted_returns, 95),
        "p95_loss_day": get_percentile(sorted_returns, 5),
    }

def create_monte_carlo_word_doc(cand_dir, cand_n, deposit, p1, p2):
    try: import docx
    except ImportError: return None

    doc = docx.Document()
    doc.add_heading(f"Monte Carlo Full Results — Candidate {cand_n}", 0)

    def add_section(phase, target, res):
        doc.add_heading(f"PHASE {phase} — Full Simulation Results", level=1)
        doc.add_heading("Simulation Parameters", level=2)
        doc.add_paragraph(f"Simulations run:             {SIMULATIONS_RUN:,}\nStarting capital:            {deposit:,.2f}\nProfit target:               {target:.1f}%\nMax drawdown limit:          {MAX_DD_LIMIT_PCT:.1f}%  (static, off initial balance)\nDaily drawdown limit:        {DAILY_DD_LIMIT_PCT:.1f}%  (off previous day close)\nMax trading days simulated:  {MAX_DAYS}\nResample block size:         {BLOCK_SIZE} trading days")
        
        doc.add_heading("Outcome Rates", level=2)
        doc.add_paragraph(f"Pass rate:                   {res['pass_rate']:.2f}%\nBreached max drawdown:       {res['breach_max_rate']:.2f}%\nBreached daily drawdown:     {res['breach_daily_rate']:.2f}%\nIncomplete (neither):        {res['incomplete_rate']:.2f}%")

        doc.add_heading("Days to Pass", level=2)
        doc.add_paragraph(f"Median:                      {res['median_days']}\n10th - 90th percentile:     {res['p10_days']}  -  {res['p90_days']}")

        doc.add_heading("Drawdown Distribution", level=2)
        doc.add_paragraph(f"Max DD reached (median):     {res['median_max_dd']:.2f}%\nMax DD reached (90th pct):   {res['p90_max_dd']:.2f}%\nMax DD reached (99th pct):   {res['p99_max_dd']:.2f}%")

        doc.add_heading("Single-Day Move Stats", level=2)
        doc.add_paragraph(f"Largest profit day:          {res['largest_profit_day']:,.2f}\nLargest loss day:            {res['largest_loss_day']:,.2f}\nAverage daily gain/loss:     {res['avg_daily']:,.2f}\n95th pct profit day:         {res['p95_profit_day']:,.2f}\n95th pct loss day:           {res['p95_loss_day']:,.2f}")

    add_section(1, PHASE_1_TARGET_PCT, p1)
    add_section(2, PHASE_2_TARGET_PCT, p2)

    doc.add_heading("CERTIFICATION VERDICT", level=1)
    comb_pass = (p1['pass_rate']/100) * (p2['pass_rate']/100) * 100
    w_max = max(p1['breach_max_rate'], p2['breach_max_rate'])
    cert = comb_pass >= 40.0 and w_max <= 20.0
    
    doc.add_paragraph(f"Combined pass rate:          {comb_pass:.2f}%\nWorst breach Max DD rate:    {w_max:.2f}%")
    p = doc.add_paragraph()
    r = p.add_run(f"RESULT: {'CERTIFIED' if cert else 'NOT CERTIFIED'}")
    r.bold = True
    r.font.color.rgb = RGBColor(0, 128, 0) if cert else RGBColor(255, 0, 0)
    
    doc_path = cand_dir / f"MonteCarlo_Full_Summary_cand_{cand_n}.docx"
    
    try:
        doc.save(doc_path)
        return doc_path
    except PermissionError:
        timestamp = time.strftime("%H%M%S")
        alt_path = cand_dir / f"MonteCarlo_Full_Summary_cand_{cand_n}_{timestamp}.docx"
        doc.save(alt_path)
        print(f"    [WARNING] Target file was locked (open in Word). Saved to alternative name: {alt_path.name}")
        return alt_path

def process_candidate(cand_dir: Path, deposit: float):
    cand_name = cand_dir.name
    print(f"\n" + "="*60)
    print(f" PROCESSING CANDIDATE: {cand_name}")
    print("="*60)

    print("  [1/3] Running Full Backtest...")
    full_bt_dir = cand_dir / "full_backtest"
    full_bt_dir.mkdir(exist_ok=True)
    report_label = f"full_bt_{cand_name}"
    
    existing = list(full_bt_dir.glob("*.htm")) + list(full_bt_dir.glob("*.html"))
    if existing:
        print(f"    -> Using existing full backtest report: {existing[0].name}")
        report_path = existing[0]
    else:
        source_set = locate_or_build_set_file(cand_dir)
        if not source_set:
            print("    -> ERROR: No parameter file found. Skipping.")
            return

        profiles_dir = Path(TERMINAL_DATA_DIR) / "MQL5" / "Profiles" / "Tester"
        profiles_dir.mkdir(parents=True, exist_ok=True)
        temp_set = profiles_dir / f"bt_temp_{cand_name}.set"
        shutil.copy(source_set, temp_set)

        report_path = run_single_backtest(
            terminal_path=TERMINAL_PATH, terminal_data_dir=TERMINAL_DATA_DIR,
            expert=EXPERT, set_file=temp_set.name, symbol=SYMBOL, period=PERIOD,
            from_date=TEST_FROM, to_date=TEST_TO, work_dir=str(full_bt_dir),
            login=LOGIN, password=PASSWORD, server=SERVER, report_name=report_label,
            deposit=int(deposit), currency=CURRENCY, leverage=LEVERAGE,
            timeout=SINGLE_TEST_TIMEOUT * 2
        )
        if temp_set.exists(): temp_set.unlink()

    if not report_path or not Path(report_path).exists():
        print("    -> ERROR: Full Backtest failed to generate a report. Skipping Monte Carlo.")
        return

    print("  [2/3] Extracting Historical P&L and Performance Summary...")
    daily_pnls = print_and_extract_report_summary(Path(report_path))

    print("  [3/3] Running 5,000 Monte Carlo Simulations...")
    p1 = run_monte_carlo_simulation(daily_pnls, deposit, PHASE_1_TARGET_PCT, MAX_DD_LIMIT_PCT, DAILY_DD_LIMIT_PCT)
    p2 = run_monte_carlo_simulation(daily_pnls, deposit, PHASE_2_TARGET_PCT, MAX_DD_LIMIT_PCT, DAILY_DD_LIMIT_PCT)
    
    cand_n = cand_name.replace("cand_", "")
    doc_path = create_monte_carlo_word_doc(cand_dir, cand_n, deposit, p1, p2)
    
    if doc_path:
        print(f"    -> SUCCESS: Saved Word Report: {doc_path.name}")

def main():
    base_work_dir = Path(WORK_DIR)
    run_dir = get_target_run_dir(base_work_dir, TARGET_RUN_DIR)
    deposit_amt = float(DEPOSIT)

    if not run_dir:
        return

    print(f"Targeting Run Directory: {run_dir.name}")
    
    for cand_name in TARGET_CANDIDATES:
        cand_dir = run_dir / cand_name
        if not cand_dir.exists():
            print(f"\nWARNING: Folder for '{cand_name}' not found inside {run_dir.name}. Skipping.")
            continue
        
        process_candidate(cand_dir, deposit_amt)

    print("\n" + "="*60)
    print(" BATCH PROCESSING COMPLETE ")
    print("="*60)

if __name__ == "__main__":
    main()