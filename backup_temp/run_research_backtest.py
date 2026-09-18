"""
run_research_backtest.py - Research Strategy Default Backtest Runner

Run a single MT5 backtest for a researched strategy and save results (charts +
Word report) to:   researched_strategies/<STRATEGY_NAME>/

Usage:
    1.  Edit the USER CONFIGURATION section below.
    2.  python run_research_backtest.py

The GUI (app.py) also calls this script as a subprocess, passing overrides
via environment variables when available.
"""

import os
import time
from pathlib import Path

# ─────────────────────────────────────────────────────────────────────────────
#  USER CONFIGURATION
# ─────────────────────────────────────────────────────────────────────────────

STRATEGY_NAME   = "My_Strategy_v1"   # Folder name inside researched_strategies/
SET_FILE        = "my_strategy.set"   # Bare filename in MT5 Profiles/Tester/
EXPERT          = "TRB V1.7.ex5"     # EA .ex5 filename
SYMBOL          = "USDJPY Dukascopy"
PERIOD          = "M15"
FROM_DATE       = "2013.01.01"
TO_DATE         = "2026.07.03"

# ─────────────────────────────────────────────────────────────────────────────
#  Load overrides from AlphaForge GUI environment variables if present
# ─────────────────────────────────────────────────────────────────────────────

def _env(key, default):
    return os.environ.get(f"AF_{key}", default)

STRATEGY_NAME = _env("STRATEGY_NAME",  STRATEGY_NAME)
SET_FILE      = _env("SET_FILE",        SET_FILE)
EXPERT        = _env("EXPERT",          EXPERT)
SYMBOL        = _env("SYMBOL",          SYMBOL)
PERIOD        = _env("PERIOD",          PERIOD)
FROM_DATE     = _env("FROM_DATE",       FROM_DATE)
TO_DATE       = _env("TO_DATE",         TO_DATE)

# ─────────────────────────────────────────────────────────────────────────────
#  Project imports  (pulled from run_optimization.py config block)
# ─────────────────────────────────────────────────────────────────────────────

try:
    from run_optimization import (
        TERMINAL_PATH, TERMINAL_DATA_DIR,
        LOGIN, PASSWORD, SERVER,
        DEPOSIT, CURRENCY, LEVERAGE,
        SINGLE_TEST_TIMEOUT, WORK_DIR,
    )
except ImportError as e:
    print(f"[ERROR] Could not import run_optimization.py: {e}")
    raise

from mt5_runner import run_single_backtest
from report_analysis import analyze

try:
    from run_full_backtest import (
        parse_mt5_html_for_deals,
        generate_equity_charts,
        generate_monthly_heatmap,
        calculate_derived_metrics,
    )
    _HAS_FULL_BT = True
except ImportError:
    _HAS_FULL_BT = False

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False

# Allow DEPOSIT override from env
RESEARCH_DEPOSIT = float(_env("DEPOSIT", str(DEPOSIT)))


# ─────────────────────────────────────────────────────────────────────────────
#  Helpers
# ─────────────────────────────────────────────────────────────────────────────

def _fmt(val, prefix="$", suffix="", decimals=2):
    if val is None:
        return "N/A"
    try:
        return f"{prefix}{float(val):,.{decimals}f}{suffix}"
    except (TypeError, ValueError):
        return str(val)


def _create_research_word_report(report_dir: Path, strategy_name: str,
                                  curated: dict, derived: dict,
                                  from_date: str, to_date: str,
                                  deposit: float):
    if not _HAS_DOCX:
        print("    [WARN] python-docx not installed — skipping Word report.")
        return None

    doc = docx.Document()

    # Title
    title = doc.add_heading(f"Research Backtest Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph()
    sub_run = sub.add_run(f"Strategy: {strategy_name}")
    sub_run.bold = True
    sub_run.font.size = Pt(14)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Period: {from_date}  →  {to_date}  |  "
                      f"Deposit: ${deposit:,.2f}  |  Generated: {time.strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(
        "This is the DEFAULT (un-optimised) backtest for the researched strategy. "
        "It establishes a baseline before running the optimization pipeline."
    )

    # Key metrics table
    doc.add_heading("Key Performance Metrics", level=1)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    hdr = tbl.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"

    d = {**curated, **derived}
    metrics_to_show = [
        ("Net Profit ($)",             d.get("Net Profit") or d.get("total_net_profit")),
        ("Net Profit (%)",             d.get("Net Profit %") or d.get("profit_pct")),
        ("Profit Factor",              d.get("Profit Factor") or d.get("profit_factor")),
        ("Sharpe Ratio",               d.get("Sharpe Ratio") or d.get("sharpe")),
        ("Max Balance Drawdown (%)",   d.get("Max Balance Drawdown (%)") or d.get("max_dd_pct")),
        ("Recovery Factor",            d.get("Recovery Factor") or d.get("recovery_factor")),
        ("Win Rate (%)",               d.get("Win Rate %") or d.get("win_rate")),
        ("Total Trades",               d.get("Total Trades") or d.get("total_trades")),
        ("Largest Profit Trade ($)",   d.get("Largest Profit Trade") or d.get("largest_profit")),
        ("Largest Loss Trade ($)",     d.get("Largest Loss Trade") or d.get("largest_loss")),
    ]
    for label, val in metrics_to_show:
        row = tbl.add_row().cells
        row[0].text = str(label)
        try:
            row[1].text = f"{float(val):,.4f}" if isinstance(val, float) else str(val or "N/A")
        except (TypeError, ValueError):
            row[1].text = str(val or "N/A")

    # Charts
    chart_files = list(report_dir.glob("chart_*.png")) + list(report_dir.glob("*.png"))
    if chart_files:
        doc.add_heading("Charts", level=1)
        for ch in chart_files[:4]:
            try:
                doc.add_paragraph(ch.stem.replace("_", " ").title()).runs[0].bold = True
                doc.add_picture(str(ch), width=Inches(6.2))
                doc.add_paragraph()
            except Exception:
                pass

    doc_path = report_dir / f"{strategy_name}_Research_Report.docx"
    try:
        doc.save(doc_path)
        return doc_path
    except PermissionError:
        alt = report_dir / f"{strategy_name}_Research_Report_{int(time.time())}.docx"
        doc.save(alt)
        return alt


# ─────────────────────────────────────────────────────────────────────────────
#  Main runner
# ─────────────────────────────────────────────────────────────────────────────

def run_research_backtest(
    strategy_name: str = STRATEGY_NAME,
    set_file: str = SET_FILE,
    expert: str = EXPERT,
    symbol: str = SYMBOL,
    period: str = PERIOD,
    from_date: str = FROM_DATE,
    to_date: str = TO_DATE,
    deposit: float = None,
):
    deposit = deposit if deposit is not None else RESEARCH_DEPOSIT

    # Output directory
    research_base = Path(WORK_DIR).parent / "researched_strategies"
    report_dir = research_base / strategy_name
    report_dir.mkdir(parents=True, exist_ok=True)

    print("\n" + "=" * 72)
    print(f"  RESEARCH BACKTEST — {strategy_name}")
    print(f"  Symbol: {symbol}  |  Period: {period}  |  Expert: {expert}")
    print(f"  Period: {from_date}  →  {to_date}")
    print(f"  Deposit: ${deposit:,.2f}  |  Output: {report_dir}")
    print("=" * 72 + "\n")

    report_name = f"{strategy_name}_default"

    report_path = run_single_backtest(
        terminal_path=TERMINAL_PATH,
        terminal_data_dir=TERMINAL_DATA_DIR,
        expert=expert,
        set_file=set_file,
        symbol=symbol,
        period=period,
        from_date=from_date,
        to_date=to_date,
        work_dir=str(report_dir),
        login=LOGIN,
        password=PASSWORD,
        server=SERVER,
        report_name=report_name,
        deposit=int(deposit),
        currency=CURRENCY,
        leverage=LEVERAGE,
        timeout=SINGLE_TEST_TIMEOUT,
    )

    print(f"\n  → Report received: {report_path.name}")

    # Parse results
    result   = analyze(report_path)
    curated  = result.get("curated_summary", {})
    derived  = {}

    # Generate equity charts
    if _HAS_FULL_BT:
        deals_df = parse_mt5_html_for_deals(report_path, deposit=float(deposit))
        if deals_df is not None and not deals_df.empty:
            chart_paths = generate_equity_charts(deals_df, float(deposit), report_dir)
            derived = calculate_derived_metrics(deals_df, float(deposit))
            print(f"  → Charts generated: {list(chart_paths.keys())}")

    # Generate Word report
    doc_path = _create_research_word_report(
        report_dir, strategy_name, curated, derived, from_date, to_date, deposit
    )

    # Print key summary
    print("\n" + "─" * 72)
    print(f"  RESULTS SUMMARY — {strategy_name}")
    print("─" * 72)
    net_p  = curated.get("Net Profit") or derived.get("total_net_profit", 0) or 0
    pct    = curated.get("Net Profit %") or derived.get("profit_pct", 0) or 0
    pf     = curated.get("Profit Factor") or derived.get("profit_factor", 0) or 0
    sharpe = curated.get("Sharpe Ratio") or derived.get("sharpe", 0) or 0
    maxdd  = curated.get("Max Balance Drawdown (%)") or derived.get("max_dd_pct", 0) or 0
    trades = curated.get("Total Trades") or derived.get("total_trades", 0) or 0
    wr     = curated.get("Win Rate %") or derived.get("win_rate", 0) or 0

    print(f"  Net Profit:     ${float(net_p):>12,.2f}  ({float(pct):.2f}%)")
    print(f"  Profit Factor:   {float(pf):>10.3f}")
    print(f"  Sharpe Ratio:    {float(sharpe):>10.3f}")
    print(f"  Max Drawdown:    {float(maxdd):>9.2f}%")
    print(f"  Total Trades:    {int(trades):>10,}")
    print(f"  Win Rate:        {float(wr):>9.2f}%")
    print("─" * 72)
    print(f"\n  Output folder:  {report_dir}")
    if doc_path:
        print(f"  Word report:    {doc_path.name}")
    print("=" * 72 + "\n")

    return {
        "report_path":  report_path,
        "report_dir":   report_dir,
        "curated":      curated,
        "derived":      derived,
        "doc_path":     doc_path,
    }


# ─────────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    run_research_backtest()
