"""
run_wf_pipeline.py - Single Candidate Walk Forward Testing Runner

Flow
----
  1. Resolves candidate directory based on TARGET_RUN_DIR and TARGET_CANDIDATE.
  2. Targets the specific candidate folder.
  3. Pre-checks for existing Walk Forward runs; if they exist, parses them and skips backtesting.
  4. Generates rolling time windows and executes sequential MT5 backtests (with auto-retry).
  5. Evaluates strict WF stability criteria (Profitability >= 60%, Max DD <= 20%, Max Single Run Profit <= 50%).
  6. Compiles an advanced 'WalkForward_Full_Summary_cand_XXX.docx' into its folder.
"""

import shutil
import json
import time
import statistics
from pathlib import Path
import pandas as pd

try:
    import docx
    from docx.shared import RGBColor, Pt
except ImportError:
    print("WARNING: python-docx not installed. Word document will not be generated.")
    print("Please run: pip install python-docx")

# -- Project imports --
from run_optimization import (
    TERMINAL_PATH,
    TERMINAL_DATA_DIR,
    EXPERT,
    SYMBOL,
    PERIOD,
    LOGIN,
    PASSWORD,
    SERVER,
    DEPOSIT,
    CURRENCY,
    LEVERAGE,
    SINGLE_TEST_TIMEOUT,
    TRAIN_FROM,
    HOLDOUT_TO,
    WF_WINDOW_MONTHS,
    WF_STEP_MONTHS,
    WORK_DIR,
)
from walk_forward import (
    generate_rolling_windows,
    build_walk_forward_table,
    compute_stability_stats,
    _print_stability_report,
)
from mt5_runner import run_single_backtest
from report_analysis import analyze


# ===========================================================================
#  USER CONFIGURATION - Set run directory & candidate to test
# ===========================================================================

# Target run folder (e.g., "run_20260911_094409")
# Set to "" or "latest" to automatically pick the most recent run
TARGET_RUN_DIR = 'run_20260911_094409'

# Target candidate folder name
TARGET_CANDIDATE = 'cand_007'

# ===========================================================================


def resolve_candidate_dir(base_work_dir: Path, target_run: str, target_cand: str) -> Path | None:
    """Locates target candidate folder inside explicit run directory or latest run."""
    if target_run and target_run.strip().lower() != "latest":
        run_dir = base_work_dir / target_run.strip()
        if not run_dir.exists():
            print(f"ERROR: Specified run directory '{target_run}' does not exist under {base_work_dir}.")
            return None
        
        # Check passed_candidates subfolder first
        passed_path = run_dir / "passed_candidates" / target_cand
        if passed_path.exists():
            return passed_path
            
        # Check direct run folder
        direct_path = run_dir / target_cand
        if direct_path.exists():
            return direct_path
            
        print(f"ERROR: Candidate '{target_cand}' not found in '{run_dir}' or '{run_dir / 'passed_candidates'}'.")
        return None
    else:
        # Search across run directories sorted newest first
        run_dirs = sorted([d for d in base_work_dir.glob("run_*") if d.is_dir()], reverse=True)
        for run_dir in run_dirs:
            passed_path = run_dir / "passed_candidates" / target_cand
            if passed_path.exists():
                return passed_path
                
            direct_path = run_dir / target_cand
            if direct_path.exists():
                return direct_path
                
        print(f"ERROR: Could not locate candidate '{target_cand}' in any recent run folders under {base_work_dir}.")
        return None


def locate_or_build_set_file(cand_dir: Path) -> Path | None:
    cand_name = cand_dir.name
    cand_n = cand_name.replace("cand_", "")
    parent_dir = cand_dir.parent

    target_set_name = f"c{cand_n}_full.set"
    target_set_path = cand_dir / target_set_name
    
    if target_set_path.exists() and target_set_path.is_file():
        print(f"  [Found Explicit Set File]: {target_set_path.name}")
        return target_set_path

    candidates_to_check = []
    for ext in ["*.set", "*.SET", "*.ini", "*.INI"]:
        candidates_to_check.extend(list(cand_dir.glob(ext)))
        candidates_to_check.extend(list(cand_dir.rglob(ext)))

    if parent_dir and parent_dir.exists():
        for pattern in [f"{cand_name}*.set", f"*{cand_n}*.set", f"{cand_name}*.SET"]:
            candidates_to_check.extend(list(parent_dir.glob(pattern)))

    unique_sets = []
    for f in candidates_to_check:
        if f.is_file() and f not in unique_sets:
            unique_sets.append(f)

    if unique_sets:
        found_file = unique_sets[0]
        print(f"  [Found Fallback Set File]: {found_file}")
        return found_file

    json_files = list(cand_dir.glob("*.json")) + list(cand_dir.rglob("*.json"))
    if json_files:
        try:
            with open(json_files[0], "r", encoding="utf-8") as f:
                params = json.load(f)
            generated_set = cand_dir / f"{cand_name}.set"
            with open(generated_set, "w", encoding="utf-16-le") as f:
                for k, v in params.items():
                    f.write(f"{k}={v}\n")
            print(f"  [Generated Set File]: Created {generated_set.name}")
            return generated_set
        except Exception:
            pass

    print(f"\n  ERROR: Could not locate ANY .set file for {cand_name}.")
    return None


def _wf_with_retries(func, label: str, max_retries: int = 2, retry_delay: int = 100):
    for attempt in range(1, max_retries + 1):
        try:
            report_path = func()
            if report_path:
                return report_path
        except Exception as exc:
            print(f"      [Attempt {attempt}/{max_retries} Exception] {label}: {exc}")

        if attempt < max_retries:
            print(f"      [Warning] MT5 timed out for {label}. Waiting {retry_delay}s to retry...")
            time.sleep(retry_delay)
        else:
            print(f"      [Error] {label} completely failed after {max_retries} attempts.")
    return None


def safe_float(val, default=0.0) -> float:
    """Helper to safely convert metric strings/numbers to float."""
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    cleaned = str(val).replace(',', '').replace('%', '').strip()
    if not cleaned or cleaned.lower() == 'nan':
        return default
    try:
        return float(cleaned)
    except ValueError:
        return default


def clean_curated(raw) -> dict:
    """Recursively flattens dicts, lists, Series, or DataFrames into a clean metric dictionary."""
    result = {}
    if isinstance(raw, dict):
        for k, v in raw.items():
            if isinstance(v, (dict, list, pd.Series, pd.DataFrame)):
                sub = clean_curated(v)
                if isinstance(sub, dict):
                    result.update(sub)
                else:
                    result[str(k)] = sub
            else:
                result[str(k)] = v
    elif isinstance(raw, list):
        for item in raw:
            if isinstance(item, dict):
                result.update(clean_curated(item))
            elif isinstance(item, tuple) and len(item) == 2:
                result[str(item[0])] = item[1]
    elif isinstance(raw, pd.DataFrame):
        if not raw.empty:
            row_dict = raw.iloc[0].to_dict()
            result.update(clean_curated(row_dict))
    elif isinstance(raw, pd.Series):
        result.update(clean_curated(raw.to_dict()))
    return result


def extract_metric(r: dict, keys: list, default=0.0) -> float:
    """Extracts a metric searching multiple possible key names across the run result dict."""
    for k in keys:
        if k in r:
            f = safe_float(r[k], None)
            if f is not None:
                return f

    for dict_key in ["curated_summary", "curated", "metrics", "summary"]:
        sub = r.get(dict_key)
        if isinstance(sub, dict):
            for k in keys:
                if k in sub:
                    f = safe_float(sub[k], None)
                    if f is not None:
                        return f
                    
    return default


def create_advanced_wf_word_doc(cand_dir: Path, wf_results: list, cand_n: str, deposit: float):
    """Generates a highly detailed Word document evaluating WF passing criteria."""
    try:
        import docx
    except ImportError:
        return None

    doc = docx.Document()
    doc.add_heading(f"Walk Forward Stability Report - Candidate {cand_n}", 0)
    
    # 1. Extract and Clean Data
    profits = []
    drawdowns = []
    pfs = []
    
    for r in wf_results:
        np = extract_metric(r, ["Net Profit", "Profit", "NetProfit", "Total Net Profit"])
        dd = extract_metric(r, ["Equity Drawdown %", "Drawdown %", "Max Drawdown %", "Equity DD %", "DD %"])
        pf = extract_metric(r, ["Profit Factor", "PF"], default=1.0)
        
        profits.append(np)
        drawdowns.append(dd)
        pfs.append(pf)
    
    # 2. Compute Global Metrics
    total_runs = len(profits)
    profitable_runs = sum(1 for p in profits if p > 0)
    profitability_pct = (profitable_runs / total_runs) * 100 if total_runs else 0
    
    total_profit = sum(profits)
    max_profit = max(profits) if profits else 0
    max_profit_pct = (max_profit / deposit) * 100 if deposit > 0 else 0
    
    max_dd = max(drawdowns) if drawdowns else 0
    avg_pf = sum(pfs) / total_runs if total_runs else 0
    mean_profit = total_profit / total_runs if total_runs else 0
    
    # CV % (Overfitting Check)
    if total_runs > 1 and mean_profit != 0:
        cv = (statistics.stdev(profits) / abs(mean_profit)) * 100
    else:
        cv = 0
        
    # 3. Passing Criteria Evaluator
    pass_profitability = profitability_pct >= 60.0
    pass_single_profit = True
    if total_profit > 0:
        pass_single_profit = (max_profit / total_profit) <= 0.50
    pass_dd = max_dd <= 20.0
    
    is_passed = pass_profitability and pass_single_profit and pass_dd
    
    # 4. Write Header & Status
    p = doc.add_paragraph()
    run_status = p.add_run(f"WF TESTING STATUS: {'PASSED VERIFICATION' if is_passed else 'FAILED VERIFICATION'}")
    run_status.bold = True
    run_status.font.size = Pt(16)
    run_status.font.color.rgb = RGBColor(0, 128, 0) if is_passed else RGBColor(255, 0, 0)
        
    # 5. Verification Checklist
    doc.add_heading("1. Verification Criteria Checklist", level=1)
    doc.add_paragraph(f"[{'X' if pass_profitability else ' '}] At least 60% profitability across all runs (Actual: {profitability_pct:.2f}%)")
    single_run_share = (max_profit / total_profit) * 100 if total_profit > 0 else 0
    doc.add_paragraph(f"[{'X' if pass_single_profit else ' '}] No single run > 50% of total profits (Highest Run Share: {single_run_share:.2f}%)")
    doc.add_paragraph(f"[{'X' if pass_dd else ' '}] No single run > 20% max drawdown (Actual Max DD: {max_dd:.2f}%)")

    # 6. Global Summary Section
    doc.add_heading("2. Overall Stability Metrics", level=1)
    doc.add_paragraph(f"• Total Runs Evaluated: {total_runs}")
    doc.add_paragraph(f"• Profitable Runs: {profitable_runs}/{total_runs} ({profitability_pct:.2f}%)")
    doc.add_paragraph(f"• Total Combined Net Profit: ${total_profit:,.2f}")
    doc.add_paragraph(f"• Max Single Run Profit: ${max_profit:,.2f} ({max_profit_pct:.2f}%)")
    doc.add_paragraph(f"• Max Drawdown Across All Runs: {max_dd:.2f}%")
    doc.add_paragraph(f"• Average Profit Factor Across Runs: {avg_pf:.2f}")
    doc.add_paragraph(f"• Coefficient of Variation (CV) [Overfitting Check]: {cv:.2f}%")
    
    # 7. Run-by-Run Table
    doc.add_heading("3. Run-by-Run Breakdown", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Window'
    hdr_cells[1].text = 'From -> To'
    hdr_cells[2].text = 'Net Profit ($)'
    hdr_cells[3].text = 'Profit (%)'
    hdr_cells[4].text = 'Max DD (%)'
    hdr_cells[5].text = 'Profit Factor'
    
    for idx, r in enumerate(wf_results):
        np = profits[idx]
        np_pct = (np / float(deposit)) * 100 if float(deposit) else 0
        dd = drawdowns[idx]
        pf = pfs[idx]
        
        from_d = r.get("from_date", r.get("from", ""))
        to_d = r.get("to_date", r.get("to", ""))
        
        row_cells = table.add_row().cells
        row_cells[0].text = f"W{idx+1:02d}"
        row_cells[1].text = f"{from_d} -> {to_d}"
        row_cells[2].text = f"${np:,.2f}"
        row_cells[3].text = f"{np_pct:.2f}%"
        row_cells[4].text = f"{dd:.2f}%"
        row_cells[5].text = f"{pf:.2f}"
        
    doc_path = cand_dir / f"WalkForward_Full_Summary_cand_{cand_n}.docx"
    
    # Safe save handling in case file is locked
    try:
        doc.save(doc_path)
        return doc_path
    except PermissionError:
        timestamp = time.strftime("%H%M%S")
        alt_path = cand_dir / f"WalkForward_Full_Summary_cand_{cand_n}_{timestamp}.docx"
        doc.save(alt_path)
        print(f"    [WARNING] File open in Word. Saved to alternative name: {alt_path.name}")
        return alt_path


def run_wf_for_candidate(cand_dir: Path, windows: list[tuple[str, str]]) -> None:
    cand_name = cand_dir.name
    cand_n = cand_name.replace("cand_", "")
    print(f"\n{'='*72}\n  Running Walk Forward Testing: Candidate {cand_n} ({cand_dir})\n{'='*72}")

    wf_dir = cand_dir / "walk_forward_runs"
    wf_dir.mkdir(exist_ok=True)
    wf_results = []

    # A. Pre-check for existing runs
    runs_already_exist = True
    for idx, (win_from, win_to) in enumerate(windows, start=1):
        window_label = f"wf_w{idx:02d}"
        run_folder = wf_dir / window_label
        report_path = None
        
        if run_folder.exists():
            for ext in ["*.htm", "*.html", "*.xml"]:
                found = list(run_folder.glob(ext))
                if found:
                    report_path = found[0]
                    break
                    
        if report_path:
            try:
                res = analyze(report_path)
                raw_curated = res.get("curated_summary", res.get("curated", res))
                curated = clean_curated(raw_curated)
                
                wf_results.append({
                    "window_idx": idx,
                    "window_index": idx,
                    "Window": idx,
                    "from_date": win_from,
                    "to_date": win_to,
                    "curated_summary": curated,
                    "curated": curated,
                    **curated,
                })
            except Exception:
                runs_already_exist = False
                break
        else:
            runs_already_exist = False
            break

    # B. Run MT5 tests if missing
    if runs_already_exist and len(wf_results) == len(windows):
        print(f"\n  >>> COMPLETE WALK FORWARD DATA ALREADY EXISTS FOR {cand_name} <<<")
        print("  >>> Skipping MT5 backtests and parsing existing reports directly. <<<")
    else:
        wf_results = []
        source_set_file = locate_or_build_set_file(cand_dir)
        if not source_set_file:
            print(f"  Aborting Walk Forward test for {cand_name}.")
            return

        profiles_tester_dir = Path(TERMINAL_DATA_DIR) / "MQL5" / "Profiles" / "Tester"
        profiles_tester_dir.mkdir(parents=True, exist_ok=True)
        temp_set_name = f"wf_temp_{cand_name}.set"
        target_set_path = profiles_tester_dir / temp_set_name
        shutil.copy(source_set_file, target_set_path)

        for idx, (win_from, win_to) in enumerate(windows, start=1):
            window_label = f"wf_w{idx:02d}"
            print(f"  --> Window {idx}/{len(windows)}: {win_from} -> {win_to}")

            def _launch():
                return run_single_backtest(
                    terminal_path=TERMINAL_PATH,
                    terminal_data_dir=TERMINAL_DATA_DIR,
                    expert=EXPERT,
                    set_file=temp_set_name,
                    symbol=SYMBOL,
                    period=PERIOD,
                    from_date=win_from,
                    to_date=win_to,
                    work_dir=str(wf_dir / window_label),
                    login=LOGIN,
                    password=PASSWORD,
                    server=SERVER,
                    report_name=window_label,
                    deposit=int(DEPOSIT),
                    currency=CURRENCY,
                    leverage=LEVERAGE,
                    timeout=SINGLE_TEST_TIMEOUT,
                )

            report = _wf_with_retries(_launch, label=f"WF {cand_name} [{window_label}]", max_retries=2, retry_delay=100)
            
            if report is None:
                print(f"      FAILED to generate report for window {idx}. Skipping window.")
                continue

            try:
                res = analyze(report)
                raw_curated = res.get("curated_summary", res.get("curated", res))
                curated = clean_curated(raw_curated)
                
                wf_results.append({
                    "window_idx": idx,
                    "window_index": idx,
                    "Window": idx,
                    "from_date": win_from,
                    "to_date": win_to,
                    "curated_summary": curated,
                    "curated": curated,
                    **curated,
                })
                print(f"      Net Profit: ${extract_metric(wf_results[-1], ['Net Profit', 'Profit']):,.2f} | Max DD: {extract_metric(wf_results[-1], ['Equity Drawdown %', 'Drawdown %'])}%")
            except Exception as exc:
                print(f"      ERROR analyzing report for window {idx}: {exc}")

        if target_set_path.exists():
            target_set_path.unlink()

    if not wf_results:
        print(f"  WARNING: No data available to generate summaries for {cand_name}.")
        return

    # C. Generate terminal stats & Word doc
    try:
        wf_table = build_walk_forward_table(wf_results)
        wf_table.to_csv(cand_dir / "walk_forward_table.csv", index=False)
    except Exception as exc:
        print(f"  [Notice] Skipping CSV build due to layout mismatch: {exc}")

    try:
        wf_stats = compute_stability_stats(wf_results)
        _print_stability_report(wf_stats)
    except Exception as exc:
        print(f"  [Notice] Skipping terminal stats printout: {exc}")

    try:
        doc_path = create_advanced_wf_word_doc(cand_dir, wf_results, cand_n, float(DEPOSIT))
        if doc_path:
            print(f"\n  SUCCESS: Advanced Walk Forward Summary Document Generated!")
            print(f"  Path: {doc_path.name}")
    except Exception as exc:
        print(f"  ERROR generating Word Doc: {exc}")


def main():
    base_work_dir = Path(WORK_DIR)
    target_cand_dir = resolve_candidate_dir(base_work_dir, TARGET_RUN_DIR, TARGET_CANDIDATE)

    print("=" * 72)
    print(f"SINGLE CANDIDATE WALK FORWARD TESTING")
    print("=" * 72)
    
    if target_cand_dir is None or not target_cand_dir.exists():
        return

    print(f"  Target Run Directory:  {target_cand_dir.parent.name}")
    print(f"  Target Candidate:      {TARGET_CANDIDATE}")
    print(f"  Full Period Range:     {TRAIN_FROM} -> {HOLDOUT_TO}")
    print(f"  Window Size:           {WF_WINDOW_MONTHS} Months")
    print(f"  Step Size:             {WF_STEP_MONTHS} Months")

    windows = generate_rolling_windows(TRAIN_FROM, HOLDOUT_TO, WF_WINDOW_MONTHS, WF_STEP_MONTHS)
    print(f"  Generated {len(windows)} rolling windows.\n")

    run_wf_for_candidate(target_cand_dir, windows)

    print("\n" + "=" * 72)
    print(f"WF PIPELINE COMPLETE FOR {TARGET_CANDIDATE}")
    print("=" * 72)


if __name__ == "__main__":
    main()